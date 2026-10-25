from pathlib import Path
from docx import Document

data_folder = Path(__file__).parent / "data"
files = sorted(data_folder.glob("*.docx"))

for file in files:
    document = Document(file)

    print("\nFILE:", file.name)
    print("Body paragraphs:", len(document.paragraphs))
    print("Tables:", len(document.tables))

    paragraphs = []
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            paragraphs.append(paragraph.text)

    print("Preview:")
    for text in paragraphs[:3]:
        print(text[:200])
        total_words = 0

for file in files:
    document = Document(file)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    word_count = len(text.split())

    print(file.name, "—", word_count, "words")
    total_words = total_words + word_count

print("Total approximate word count:", total_words)